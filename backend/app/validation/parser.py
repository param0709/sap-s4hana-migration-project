"""Read uploaded CSV/XLSX files without losing source-column identity.

Headers and row values are read positionally. This avoids pandas' automatic
duplicate-header renaming, which would hide FR-009 findings and could overwrite
source values when rows are converted to JSON objects.
"""
import csv
import io
import math
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook

from app.validation.errors import FileRejectedError

_CSV_ENCODINGS = ("utf-8-sig", "utf-8", "latin-1")


def _json_safe(value: Any) -> Any:
    """Convert one source cell to a value accepted by JSON/JSONB."""
    if value is None:
        return None
    if isinstance(value, str):
        return None if value == "" else value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return value
    if isinstance(value, Decimal):
        number = float(value)
        return int(number) if number.is_integer() else number
    if isinstance(value, float):
        return None if math.isnan(value) or math.isinf(value) else value
    if isinstance(value, int):
        return value

    # Defensive support for scalar types supplied by third-party readers.
    item = getattr(value, "item", None)
    if callable(item):
        try:
            return _json_safe(item())
        except (TypeError, ValueError):
            pass
    return str(value)


def _is_blank_cell(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _is_blank_row(row: Sequence[Any]) -> bool:
    return all(_is_blank_cell(value) for value in row)


def build_record_keys(headers: list[str]) -> list[str]:
    """Create ordered, normalised and duplicate-safe JSON keys."""
    counts: dict[str, int] = {}
    keys: list[str] = []

    for position, header in enumerate(headers, start=1):
        name = str(header).strip().upper()
        if not name:
            keys.append(f"__BLANK_COLUMN_{position}")
            continue

        occurrence = counts.get(name, 0) + 1
        counts[name] = occurrence
        keys.append(name if occurrence == 1 else f"{name}__DUPLICATE_{occurrence}")

    return keys


def _normalise_rows(
    rows: Iterable[Sequence[Any]],
    keys: list[str],
    file_name: str,
) -> tuple[list[dict[str, Any]], list[int]]:
    """Validate row width and map nonblank rows onto duplicate-safe keys.

    Short rows are padded with nulls. Empty cells beyond the header width are
    harmless, but a nonblank extra cell is rejected because silently dropping
    it would violate source-data preservation.
    """
    width = len(keys)
    records: list[dict[str, Any]] = []
    source_row_numbers: list[int] = []

    # Row 1 is the source header. Starting at 2 preserves the row number a
    # consultant sees in Excel or a CSV editor, including gaps from blank rows.
    for source_row_number, source_row in enumerate(rows, start=2):
        row = list(source_row)
        if _is_blank_row(row):
            continue

        extras = row[width:]
        if any(not _is_blank_cell(value) for value in extras):
            raise FileRejectedError(
                "MALFORMED_ROW",
                f"'{file_name}' contains a data row wider than its header.",
                {
                    "source_row_number": source_row_number,
                    "expected_column_count": width,
                    "actual_column_count": len(row),
                },
            )

        record = {
            key: _json_safe(row[index]) if index < len(row) else None
            for index, key in enumerate(keys)
        }
        records.append(record)
        source_row_numbers.append(source_row_number)

    return records, source_row_numbers


@dataclass(frozen=True)
class ParsedFile:
    """Raw ordered headers plus one JSON-safe object per nonblank data row."""

    headers: list[str]
    row_count: int
    rows: list[dict[str, Any]] = field(default_factory=list)
    source_row_numbers: list[int] = field(default_factory=list)

    @property
    def column_count(self) -> int:
        return len(self.headers)


def _decode_csv(content: bytes, file_name: str) -> str:
    for encoding in _CSV_ENCODINGS:
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FileRejectedError(
        "CORRUPTED_FILE",
        f"'{file_name}' uses an unsupported text encoding.",
    )


def _read_csv(content: bytes, file_name: str) -> ParsedFile:
    text = _decode_csv(content, file_name)
    try:
        all_rows = list(csv.reader(io.StringIO(text, newline=""), strict=True))
    except csv.Error as exc:
        raise FileRejectedError(
            "CORRUPTED_FILE",
            f"'{file_name}' could not be read as a CSV file.",
            {"reason": str(exc)[:300]},
        ) from exc

    if not all_rows:
        raise FileRejectedError(
            "EMPTY_FILE",
            f"'{file_name}' contains no rows or columns.",
        )

    headers = [str(value) for value in all_rows[0]]
    rows, source_row_numbers = _normalise_rows(
        all_rows[1:], build_record_keys(headers), file_name
    )
    return ParsedFile(
        headers=headers,
        row_count=len(rows),
        rows=rows,
        source_row_numbers=source_row_numbers,
    )


def _read_excel(content: bytes, file_name: str) -> ParsedFile:
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        all_rows = list(sheet.iter_rows(values_only=True))
    except Exception as exc:
        raise FileRejectedError(
            "CORRUPTED_FILE",
            f"'{file_name}' could not be read as an Excel workbook.",
            {"reason": str(exc)[:300]},
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()

    if not all_rows:
        return ParsedFile(headers=[], row_count=0, rows=[])

    header_cells = list(all_rows[0])
    # openpyxl's used range includes columns populated only in later rows. Those
    # trailing None cells are not declared headers and are removed here; any
    # nonblank data below them is subsequently rejected as a malformed row.
    while header_cells and header_cells[-1] is None:
        header_cells.pop()

    headers = ["" if value is None else str(value) for value in header_cells]
    rows, source_row_numbers = _normalise_rows(
        all_rows[1:], build_record_keys(headers), file_name
    )
    return ParsedFile(
        headers=headers,
        row_count=len(rows),
        rows=rows,
        source_row_numbers=source_row_numbers,
    )


def read_tabular_file(content: bytes, file_name: str, extension: str) -> ParsedFile:
    """Parse a supported upload and require headers plus at least one data row."""
    parsed = (
        _read_csv(content, file_name)
        if extension == ".csv"
        else _read_excel(content, file_name)
    )

    if parsed.column_count == 0:
        raise FileRejectedError("EMPTY_FILE", f"'{file_name}' has no columns.")
    if parsed.row_count == 0:
        raise FileRejectedError(
            "NO_DATA_ROWS",
            f"'{file_name}' has column headers but no data rows.",
            {"column_count": parsed.column_count},
        )
    return parsed
